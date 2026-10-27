#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""Convert a small markdown subset into a styled academic .docx.

Supported: # / ## / ### headings, paragraphs, **bold**, *italic*, `code`,
- bullets, 1. numbered items, > quote, | tables |, ```code blocks```.

Usage:  python md2docx.py input.md output.docx [--subtitle "..."] [--author "..."]
"""
import argparse
import re
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BODY_CN = "宋体"
HEAD_CN = "黑体"
LATIN = "Times New Roman"
BODY_SIZE = 12          # 小四
LINE_SPACING = 1.5


def _set_run_font(run, cn=BODY_CN, size=BODY_SIZE, bold=None, italic=None, color=None):
    run.font.name = LATIN
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), cn)
    rfonts.set(qn("w:ascii"), LATIN)
    rfonts.set(qn("w:hAnsi"), LATIN)


INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\*[^*]+?\*)")


def add_inline(paragraph, text, size=BODY_SIZE, cn=BODY_CN):
    for chunk in INLINE.split(text):
        if not chunk:
            continue
        if chunk.startswith("**") and chunk.endswith("**") and len(chunk) > 4:
            _set_run_font(paragraph.add_run(chunk[2:-2]), cn=cn, size=size, bold=True)
        elif chunk.startswith("`") and chunk.endswith("`") and len(chunk) > 2:
            _set_run_font(paragraph.add_run(chunk[1:-1]), cn=cn, size=size - 1)
        elif chunk.startswith("*") and chunk.endswith("*") and len(chunk) > 2:
            _set_run_font(paragraph.add_run(chunk[1:-1]), cn=cn, size=size, italic=True)
        else:
            _set_run_font(paragraph.add_run(chunk), cn=cn, size=size)


def para(doc, text, size=BODY_SIZE, cn=BODY_CN, indent=True, space_after=6,
         align=None, bold=False):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(space_after)
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = LINE_SPACING
    if indent:
        pf.first_line_indent = Pt(size * 2)
    if align is not None:
        p.alignment = align
    elif indent:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if bold:
        _set_run_font(p.add_run(text), cn=cn, size=size, bold=True)
    else:
        add_inline(p, text, size=size, cn=cn)
    return p


def heading(doc, text, level):
    if level == 1:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(14)
        _set_run_font(p.add_run(text), cn=HEAD_CN, size=18, bold=True)
        return p
    size = {2: 14, 3: 12.5}.get(level, 12)
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(12 if level == 2 else 8)
    pf.space_after = Pt(6)
    pf.keep_with_next = True
    _set_run_font(p.add_run(text), cn=HEAD_CN, size=size, bold=True)
    return p


def bullets(doc, items, numbered=False):
    for i, item in enumerate(items, 1):
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.left_indent = Cm(0.9)
        pf.first_line_indent = Cm(-0.55) if numbered else Cm(-0.35)
        pf.space_before = Pt(0)
        pf.space_after = Pt(2)
        pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        pf.line_spacing = 1.35
        marker = f"{i}. " if numbered else "• "
        _set_run_font(p.add_run(marker), size=BODY_SIZE)
        add_inline(p, item)


def add_table(doc, rows):
    rows = [r for r in rows if any(c.strip() for c in r)]
    if not rows:
        return
    ncol = max(len(r) for r in rows)
    table = doc.add_table(rows=0, cols=ncol)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for ri, row in enumerate(rows):
        trow = table.add_row()
        if ri == 0:
            trpr = trow._tr.get_or_add_trPr()
            trpr.append(OxmlElement("w:tblHeader"))
        cells = trow.cells
        for ci in range(ncol):
            text = row[ci].strip() if ci < len(row) else ""
            cell = cells[ci]
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            if ri == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_inline(p, text, size=10.5)
            if ri == 0:
                for run in p.runs:
                    run.bold = True
    para(doc, "", indent=False, space_after=4)


def split_row(line):
    return [c for c in line.strip().strip("|").split("|")]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("src")
    ap.add_argument("dst")
    ap.add_argument("--subtitle", default="")
    ap.add_argument("--author", default="")
    args = ap.parse_args()

    with open(args.src, "r", encoding="utf-8") as f:
        lines = f.read().replace("\r\n", "\n").split("\n")

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.6)
    sec.top_margin = sec.bottom_margin = Cm(2.5)
    normal = doc.styles["Normal"]
    normal.font.name = LATIN
    normal.font.size = Pt(BODY_SIZE)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_CN)

    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run_font(footer.add_run("第 "), size=9)
    fld = footer.add_run()
    _set_run_font(fld, size=9)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    fld._r.append(begin)
    fld._r.append(instr)
    fld._r.append(end)
    _set_run_font(footer.add_run(" 页"), size=9)

    i, n = 0, len(lines)
    pending = []          # consecutive bullet/numbered lines
    pending_num = False
    table_buf = []

    def flush_bullets():
        nonlocal pending
        if pending:
            bullets(doc, pending, numbered=pending_num)
            pending = []

    def flush_table():
        nonlocal table_buf
        if table_buf:
            add_table(doc, table_buf)
            table_buf = []

    while i < n:
        raw = lines[i]
        line = raw.rstrip()
        stripped = line.strip()

        if stripped.startswith("|") and stripped.endswith("|"):
            flush_bullets()
            cells = split_row(stripped)
            if not all(re.fullmatch(r":?-{2,}:?", c.strip()) for c in cells if c.strip()):
                table_buf.append(cells)
            i += 1
            continue
        flush_table()

        if not stripped:
            flush_bullets()
            i += 1
            continue

        if stripped.startswith("```"):
            flush_bullets()
            i += 1
            block = []
            while i < n and not lines[i].strip().startswith("```"):
                block.append(lines[i])
                i += 1
            i += 1
            for b in block:
                p = doc.add_paragraph()
                pf = p.paragraph_format
                pf.left_indent = Cm(0.6)
                pf.space_after = Pt(0)
                pf.line_spacing = 1.0
                _set_run_font(p.add_run(b), size=9.5)
            continue

        m = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if m:
            flush_bullets()
            heading(doc, m.group(2).strip(), len(m.group(1)))
            if len(m.group(1)) == 1 and (args.subtitle or args.author):
                for extra in (args.subtitle, args.author):
                    if extra:
                        p = doc.add_paragraph()
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        p.paragraph_format.space_after = Pt(4)
                        _set_run_font(p.add_run(extra), size=11)
            i += 1
            continue

        if stripped.startswith(">"):
            flush_bullets()
            p = doc.add_paragraph()
            pf = p.paragraph_format
            pf.left_indent = Cm(0.8)
            pf.space_after = Pt(4)
            add_inline(p, stripped[1:].strip(), size=BODY_SIZE - 0.5)
            i += 1
            continue

        m = re.match(r"^[-*+]\s+(.*)$", stripped)
        if m:
            if pending and pending_num:
                flush_bullets()
            pending_num = False
            pending.append(m.group(1).strip())
            i += 1
            continue

        m = re.match(r"^(\d+)[.)]\s+(.*)$", stripped)
        if m:
            if pending and not pending_num:
                flush_bullets()
            pending_num = True
            pending.append(m.group(2).strip())
            i += 1
            continue

        if stripped in ("---", "***", "___"):
            flush_bullets()
            i += 1
            continue

        flush_bullets()
        para(doc, stripped)
        i += 1

    flush_bullets()
    flush_table()
    doc.save(args.dst)

    doc2 = Document(args.dst)
    words = sum(len(p.text) for p in doc2.paragraphs)
    print(f"saved: {args.dst}")
    print(f"paragraphs={len(doc2.paragraphs)} tables={len(doc2.tables)} chars={words}")


if __name__ == "__main__":
    sys.exit(main())
